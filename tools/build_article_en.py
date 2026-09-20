from __future__ import annotations

import json
import re
import sys
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from optimml.experiment import write_plot


SOURCE = ROOT / "article" / "SUMMA2026_RU_gradient_geometry.docx"
OUTPUT = ROOT / "article" / "SUMMA2026_EN_gradient_geometry.docx"
RESULTS = ROOT / "results" / "results.json"
FIGURE = ROOT / "results" / "gradient_geometry_en.png"


TRANSLATIONS = [
    "Gradient Geometry and LoRA Adaptation in Preference Learning",
    "Petr Vladimirovich Nikitin, Candidate of Pedagogical Sciences, Associate Professor",
    "Associate Professor, Department of Artificial Intelligence",
    "Financial University under the Government of the Russian Federation",
    "Moscow, Russia; pvnikitin@fa.ru; ORCID: 0000-0001-8866-5610",
    None,
    None,
    "INTRODUCTION",
    "Learning from human preferences has become a central tool for aligning language-model behavior with user intent [1], [2]. Direct Preference Optimization (DPO) reduces policy optimization to logistic classification of preferred and rejected response pairs [3]. Later methods, including SimPO [4] and ORPO [5], modify score construction, length normalization, or the relationship to supervised fine-tuning. Despite these differences, many methods include a temperature parameter β that scales the preference margin.",
    "For a single pair, the multiplier β does not change the gradient ray because it only multiplies the margin gradient by a positive scalar. A practical optimizer step, however, is computed from a mini-batch of nonparallel contributions. Changing β redistributes their weights and can therefore rotate the aggregate gradient. Pairwise loss also depends only on the score difference, whereas pointwise loss separately constrains the scores of the preferred and rejected responses. This distinction is related to the classical Bradley-Terry model [6] and RankNet pairwise learning to rank [7].",
    "This study tests both effects on real human-preference data in a fully reproducible setting. Unlike synthetic checks reported in the original studies, our experiment uses an HH-RLHF sample and publishes all configurations, run-level results, checksums, and derivative tests. We ask how strongly β rotates the mini-batch gradient, how weight concentration changes, and whether pairwise and pointwise losses differ under the same representation and optimization budget.",
    "Parameter-efficient fine-tuning extends this setting. LoRA restricts the weight update to a low-rank subspace [13], QLoRA combines this approach with quantization [14], and LoRA+, DoRA, and PiSSA refine learning rates, weight decomposition, and adapter initialization [15]-[17]. Modern reward-model evaluation demonstrates the need for external and robust validation [18], while SePO concentrates the preference signal on informative tokens [19]. The second experiment therefore tests our geometric hypothesis directly in the LoRA parameter space.",
    "The practical relevance of this work concerns systems that must select the best of several generated responses, including contact centers, financial compliance, educational feedback, health information, industrial instructions, public services, and software review. In these settings, a reward model should serve as an auxiliary filter rather than an autonomous source of legally, clinically, or technologically significant decisions. The cross-sector NIST profile calls for context-specific assessment of generative AI risks [22], while OECD reports document both broader deployment and risks related to opacity, bias, and overreliance [23], [24].",
    "METHOD",
    "Loss Functions and Geometry",
    "For a response pair, let xᵢ=φ(yᵢʷ)−φ(yᵢˡ) and define the linear scorer sθ(y)=θᵀφ(y). The pairwise logistic loss is",
    "Its gradient is a weighted sum of the vectors xᵢ:",
    "The weights wᵢ(β)=σ(−βθᵀxᵢ) depend on the magnitude and sign of the margin. Thus, for nonparallel xᵢ, changing β generally changes the direction of the sum. Weight concentration was measured using the normalized effective sample size",
    "The pointwise loss classifies the two responses separately:",
    "The pairwise formulation is invariant to a common shift in the two scores, whereas the pointwise formulation is not. Therefore, even when the response ordering is identical, the loss functions impose different curvature and different constraints on θ.",
    "Data and Representation",
    "We used the open Anthropic HH-RLHF dataset [2], [11]. Nonoverlapping API pages were selected deterministically from the 160,800-row training split. After pairs with identical chosen and rejected texts were removed, 2,000 valid pairs remained. Sampling used seed 20260918. The SHA-256 hash of the canonical JSONL is e7a1ba050309838772fc8eca5506f601b91b556b7ab3fcddd3f77ee591e096bd.",
    "Texts were converted to 4,096-dimensional signed feature-hashing vectors using word unigrams and bigrams with L2 normalization. This method requires no vocabulary and provides a reproducible memory bound [8]. The representation was deliberately fixed because the experiment measures the geometry of preference objectives rather than the quality of a particular Transformer.",
    "The data were divided into 1,400 training, 300 validation, and 300 test pairs without overlap. For β∈{0.1, 0.3, 1, 3, 10}, both models were trained with Adam [9] for 800 steps using batch size 128, learning rate 0.03, L2=10⁻³, and seeds 11, 29, 47, 83, and 101. For the geometric diagnostic, a reference pairwise model was trained at β=1 for 1,000 steps, and gradients were computed on the unchanged validation split.",
    "Transformer and LoRA",
    "The second experiment used a pretrained English BERT-tiny model with two Transformer layers and hidden size 128. Texts were truncated from the left to 192 tokens to preserve the end of each response. A scalar reward head was applied to the [CLS] representation. The control model trained only the head. The LoRA model additionally trained the query and value matrices in every attention layer. For an original matrix W, the low-rank update was defined as",
    "We used r=8, α=16, dropout 0.05, β=0.3, AdamW, a learning rate of 5·10⁻⁴ for the adapters and 10⁻³ for the head, weight decay 0.01, batch size 16, and seeds 11, 29, and 47. Training ran for four epochs, and the checkpoint was selected exclusively by the minimum validation pairwise loss. The test split was used once after selection. LoRA trained 8,321 of 4,394,241 parameters. For the run with the best validation loss, gradients of all trainable parameters were computed on a fixed batch of 64 validation pairs at five β values.",
    "The 1.1B Model and External Evaluation",
    "To test the 1-3 billion parameter scale, we used TinyLlama-1.1B-Chat-v1.0 [21] as a Llama-based reward model with a scalar classification head. Rank-4 LoRA with α=8 and dropout 0.05 was applied to q_proj and v_proj. The experiment trained 565,248 of 1,035,079,680 parameters, or 0.0546%. Because of the CPU budget, we ran a preregistered pilot with 24 steps on 24 HH training pairs, β=0.3, and a sequence length of 64 tokens. Before and after adaptation, the model was evaluated on 48 nonoverlapping HH pairs and 92 RewardBench examples, with four examples from each of 23 subsets. This sample is a stratified diagnostic set and not an official leaderboard aggregate [18].",
    "Checks and Metrics",
    "Analytical gradients of both losses were checked by central finite differences. The absolute directional-derivative error was 6.65·10⁻¹³ for the pairwise loss and 8.11·10⁻¹¹ for the pointwise loss. We also verified ESS bounds, reference-vector angles, split separation, hashing repeatability, and data checksums. The metrics included angle, norm, ESS, saturation, ranking accuracy, log loss, Brier score, and ECE [10]. Accuracy intervals were estimated by bootstrap [12], and the change on identical RewardBench examples was additionally tested with a two-sided exact sign test.",
    "RESULTS",
    "GRADIENT GEOMETRY ON THE VALIDATION SET",
    "Table I and Fig. 1 show a nonmonotonic rotation relative to the β=1 reference. The angle was 17.22° at β=0.1, 20.89° at β=3, and 34.71° at β=10. At the same time, ESS/n decreased from 0.9998 to 0.7141. At β=10, 21.67% of the pairs already had saturated weights. Thus, changing β cannot be replaced by changing the learning rate because both the direction and the effective composition of the mini-batch change with the scale.",
    "Rotation of the mini-batch gradient and reduction in normalized ESS as β changes.",
    "TEST RESULTS ACROSS FIVE RUNS",
    "Both models achieved their best result at β=0.3. Pairwise loss reached an accuracy of 0.5627±0.0072, while pointwise loss reached 0.5433±0.0129. The mean advantage of the pairwise model across matched seeds was 1.93 percentage points, although this conclusion is limited to one sample and one feature class. At β=10, accuracy decreased to 0.4800 and 0.4833, while log loss increased to 1.069 and 1.218, respectively. Therefore, a larger gradient norm at high β did not improve ranking quality.",
    "Probability calibration remained weak in all settings. The minimum Brier score was 0.2484 for the pairwise model and 0.2492 for the pointwise model. This result is expected for a compact linear scorer and shows that ranking accuracy and confidence should be analyzed separately.",
    "LORA ADAPTATION RESULTS",
    "COMPARISON OF TRANSFORMER REWARD MODELS",
    "LoRA did not improve mean ranking accuracy: 0.5333±0.0200 compared with 0.5467±0.0058 for the control. Mean pairwise loss was also nearly identical at 0.6872 and 0.6865. All three LoRA runs selected the first epoch by validation loss, whereas the control selected epoch 1.67 on average. This pattern indicates early adapter overfitting and does not support a LoRA advantage on 1,400 training pairs. The best individual LoRA seed reached 0.5533, but this single run was not used as the main result.",
    "GRADIENT GEOMETRY IN THE LORA PARAMETER SPACE",
    "The geometric effect was reproduced in the trainable neural subspace. The angle to the β=1 gradient was 22.99° at β=0.1, 19.44° at β=3, and 30.39° at β=10. Meanwhile, ESS/n decreased to 0.6199 and the saturated-weight share increased to 0.3125. Temperature scaling therefore changes the update direction in both the linear surrogate and the LoRA adapter.",
    "EXTERNAL TINYLLAMA 1.1B PILOT",
    "ACCURACY ON THE STRATIFIED REWARDBENCH SAMPLE, %",
    "After 24 LoRA steps, TinyLlama micro accuracy on 92 external examples changed from 42.39% to 45.65%, an increase of 3.26 percentage points. Eight examples improved, five became worse, and 79 remained unchanged. However, the 95% bootstrap interval for the paired change was −4.35 to 10.87 percentage points, and the two-sided exact sign test gave p=0.581. Accuracy on 48 held-out HH pairs decreased from 47.92% to 45.83%. Chat, Chat Hard, and Safety improved, while Reasoning decreased from 50.0% to 46.4%. The pilot therefore demonstrates a reproducible transfer protocol but not a statistically supported improvement.",
    "DISCUSSION",
    "The observed relationship between angle and ESS agrees with the theoretical mechanism. Increasing β suppresses the contribution of pairs with a positive margin and concentrates the update on difficult or misranked examples. At small β, the weights are almost equal, but the direction still differs from β=1 because the reference model already creates heterogeneous margins. Angle symmetry around β=1 is not expected because logistic reweighting is nonlinear.",
    "The loss comparison shows that the distinction between difference-based and absolute identification persists on real pairs. The LoRA experiment transfers the analysis from fixed features to internal Transformer representations, but the absence of an average accuracy gain limits the practical conclusion. The result agrees with the KTO view that the appropriate loss depends on assumptions about human utility [20]. Adapter parameterization alone does not resolve a mismatch between a small sample and the target metric.",
    "Sector-specific deployment requires measuring the critical error for each setting. For customer support and e-commerce, this error concerns factual usefulness and incorrect escalation. For finance, it concerns compliance and confidentiality. Education requires correct reasoning, while health care and industry require safe refusal and mandatory expert review. Public-sector deployment requires traceability, equal treatment, and data protection. The mixed RewardBench result shows why a mean score should not be the sole deployment criterion.",
    "The study has several limitations: one English training dataset, only three BERT seeds, one seed with 24 pairs and short sequences in the 1.1B CPU pilot, 92 external examples instead of the full RewardBench, and no sector-specific annotations. Future work should run the full benchmark, increase the number of seeds and pairs, compare LoRA, LoRA+, and PiSSA on 1-3B models, add M-RewardBench and a Russian-language corpus, and preregister success criteria for every category.",
    "CONCLUSION",
    "The experiment confirmed that β is a geometric hyperparameter. The angle reached 34.71° in the linear model and 30.39° in the LoRA parameter space. BERT-tiny adaptation produced no mean advantage, while the TinyLlama 1.1B CPU pilot changed external micro accuracy by +3.26 percentage points without statistical significance and reduced held-out HH accuracy. Model scale and a small trainable-parameter share therefore do not replace external validation. In practice, β, the loss function, the LoRA configuration, and the sector-specific acceptance criterion should be selected jointly while retaining category-level metrics and human-in-the-loop oversight.",
    "DATA AND CODE AVAILABILITY",
    "Code, configurations, run-level results, checksums, and both language versions of the article are available at https://github.com/PetrNikitin20/optimML. Source pairs are downloaded from Anthropic HH-RLHF, and raw texts are not included in the repository.",
    "REFERENCES",
]


TABLE_TRANSLATIONS = {
    "Угол, °": "Angle, °",
    "Насыщение": "Saturation",
    "Потеря": "Loss",
    "Точность": "Accuracy",
    "парная": "pairwise",
    "поэлем.": "pointwise",
    "Метод": "Method",
    "Обуч. парам.": "Trainable",
    "Доля, %": "Share, %",
    "только head": "head only",
    "Категория": "Category",
    "1,1B до": "1.1B before",
    "1,1B после": "1.1B after",
}


def replace_paragraph(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def build() -> None:
    data = json.loads(RESULTS.read_text(encoding="utf-8"))
    write_plot(data["geometry"], FIGURE, language="en")

    doc = Document(SOURCE)
    paragraphs = [p for p in doc.paragraphs if p.text.strip()]
    if len(paragraphs) != 80 or len(TRANSLATIONS) != 56:
        raise RuntimeError(f"Unexpected source structure: {len(paragraphs)} paragraphs")

    for index, translation in enumerate(TRANSLATIONS):
        if translation is not None:
            replace_paragraph(paragraphs[index], translation)

    paragraphs[5].runs[0].text = "Abstract: "
    paragraphs[5].runs[1].text = (
        "This study examines how the temperature parameter β affects mini-batch gradient geometry and low-rank adaptation of a reward model. The experiment uses 2,000 real Anthropic HH-RLHF preference pairs. In the linear setting, changing β from 0.1 to 10 produced a gradient angle of 34.71° and reduced ESS/n to 0.7141. The geometric effect was reproduced in the BERT-tiny LoRA parameter space. An external CPU pilot with TinyLlama 1.1B trained 0.0546% of the parameters and was evaluated on a stratified sample of 92 RewardBench examples. Micro accuracy changed from 42.39% to 45.65%, but the 95% bootstrap interval for the change included zero (−4.35 to 10.87 percentage points, p=0.581). Accuracy on 48 held-out HH pairs decreased from 47.92% to 45.83%. The result confirms the computational feasibility of LoRA at the billion-parameter scale but demonstrates the need for category-level external evaluation."
    )
    paragraphs[6].runs[0].text = "Keywords: "
    paragraphs[6].runs[1].text = "preference optimization; gradient geometry; LoRA; reward model; RewardBench; TinyLlama"

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if run.text in TABLE_TRANSLATIONS:
                            run.text = TABLE_TRANSLATIONS[run.text]

    image_replaced = False
    for shape in doc.inline_shapes:
        blips = shape._inline.xpath(".//a:blip")
        if not blips:
            continue
        relationship_id = blips[0].get(qn("r:embed"))
        part = doc.part.related_parts[relationship_id]
        if part.content_type.startswith("image/"):
            part._blob = FIGURE.read_bytes()
            shape._inline.docPr.set("descr", "Two plots showing the gradient angle and normalized effective sample size as beta changes")
            shape._inline.docPr.set("title", "Mini-batch gradient geometry on HH-RLHF")
            image_replaced = True
            break
    if not image_replaced:
        raise RuntimeError("The article figure was not found")

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if "—" in run.text:
                run.text = run.text.replace("—", "–")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if "—" in run.text:
                            run.text = run.text.replace("—", "–")

    doc.core_properties.title = "Gradient Geometry and LoRA Adaptation in Preference Learning"
    doc.core_properties.author = "Petr Vladimirovich Nikitin"
    doc.core_properties.last_modified_by = "Petr Vladimirovich Nikitin"
    doc.core_properties.subject = "English manuscript for SUMMA 2026"
    doc.core_properties.keywords = "preference optimization, gradient geometry, LoRA, reward model, RewardBench, TinyLlama"
    doc.save(OUTPUT)

    with ZipFile(OUTPUT) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    if "—" in document_xml:
        raise RuntimeError("Em dash remains in the English document")
    visible_text = " ".join(p.text for p in Document(OUTPUT).paragraphs)
    if re.search(r"[А-Яа-яЁё]", visible_text):
        raise RuntimeError("Cyrillic text remains in the English document")
    print(OUTPUT)


if __name__ == "__main__":
    build()
