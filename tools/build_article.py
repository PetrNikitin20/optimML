from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "article" / "template" / "conference-template-a4-normalized.docx"
RESULTS = ROOT / "results" / "results.json"
LORA_RESULTS = ROOT / "results" / "lora_results.json"
EXTERNAL_RESULTS = ROOT / "results" / "external_benchmark_results.json"
FIGURE = ROOT / "results" / "gradient_geometry.png"
OUTPUT = ROOT / "article" / "SUMMA2026_RU_gradient_geometry.docx"


def clear_document_body(doc: Document):
    body = doc._element.body
    final_sect = body.sectPr
    for child in list(body):
        if child is not final_sect:
            body.remove(child)


def set_columns(section, count: int, space_dxa: int = 360):
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), str(space_dxa))
    for child in list(cols):
        cols.remove(child)


def set_font(run, size=None, bold=None, italic=None):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_body(doc, text):
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.keep_together = False
    set_font(p.add_run(text), size=10)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), size=10, italic=(level == 2))
    return p


def add_native_equation(doc, formula, number):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [4300, 500]
    for i, cell in enumerate(table.rows[0].cells):
        cell.width = Inches(widths[i] / 1440)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_w = tc_pr.find(qn("w:tcW"))
        if tc_w is None:
            tc_w = OxmlElement("w:tcW")
            tc_pr.append(tc_w)
        tc_w.set(qn("w:w"), str(widths[i]))
        tc_w.set(qn("w:type"), "dxa")
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "nil")
            borders.append(el)
        tc_pr.append(borders)
    p = table.cell(0, 0).paragraphs[0]
    p.style = doc.styles["equation"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    math_para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = formula
    math_run.append(math_text)
    math.append(math_run)
    math_para.append(math)
    p._p.append(math_para)
    pn = table.cell(0, 1).paragraphs[0]
    pn.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(pn.add_run(f"({number})"), size=10)
    return table


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def cell_margins(cell, value=50):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for edge in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def set_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is not None:
        tbl_pr.remove(borders)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "A6A6A6")
        borders.append(node)
    tbl_pr.append(borders)


def add_table(doc, caption, headers, rows, widths):
    cap = doc.add_paragraph(style="table head")
    cap.paragraph_format.keep_with_next = True
    set_font(cap.add_run(caption), size=8)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_borders(table)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header_marker = OxmlElement("w:tblHeader")
    header_marker.set(qn("w:val"), "true")
    tr_pr.append(header_marker)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade(cell, "D9E2F3")
        cell.width = Inches(widths[idx] / 1440)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(text), size=7.5, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cells[idx].width = Inches(widths[idx] / 1440)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cells[idx])
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 0.9
            set_font(p.add_run(str(text)), size=7.5)
    return table


def add_reference(doc, number, text):
    p = doc.add_paragraph(style="references")
    p.paragraph_format.keep_together = True
    set_font(p.add_run(text), size=8)


def fmt3(x):
    return f"{x:.3f}"


def build():
    data = json.loads(RESULTS.read_text(encoding="utf-8"))
    lora_data = json.loads(LORA_RESULTS.read_text(encoding="utf-8"))
    external = json.loads(EXTERNAL_RESULTS.read_text(encoding="utf-8"))
    doc = Document(TEMPLATE)
    clear_document_body(doc)

    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.top_margin = Inches(0.38)
    section.bottom_margin = Inches(1.0)
    set_columns(section, 1, 360)

    title = doc.add_paragraph(style="paper title")
    set_font(title.add_run("Геометрия градиента и LoRA-адаптация при обучении по предпочтениям"), size=24)

    author = doc.add_paragraph(style="Author")
    author.paragraph_format.space_before = Pt(6)
    set_font(author.add_run("Алексей П. Малахов"), size=11)
    affiliation = doc.add_paragraph(style="Author")
    affiliation.paragraph_format.space_before = Pt(0)
    set_font(affiliation.add_run("Департамент искусственного интеллекта, факультет информационных технологий и анализа больших данных"), size=10, italic=True)
    organization = doc.add_paragraph(style="Author")
    organization.paragraph_format.space_before = Pt(0)
    set_font(organization.add_run("Финансовый университет при Правительстве Российской Федерации, Москва, Россия"), size=10, italic=True)
    contact = doc.add_paragraph(style="Author")
    contact.paragraph_format.space_before = Pt(0)
    contact.paragraph_format.space_after = Pt(8)
    set_font(contact.add_run("[указать электронную почту или ORCID перед подачей]"), size=9)

    body_section = doc.add_section(WD_SECTION.CONTINUOUS)
    body_section.left_margin = Inches(0.62)
    body_section.right_margin = Inches(0.62)
    body_section.top_margin = Inches(0.75)
    body_section.bottom_margin = Inches(1.0)
    set_columns(body_section, 2, 360)

    abstract = doc.add_paragraph(style="Abstract")
    abstract.paragraph_format.first_line_indent = Inches(0)
    r = abstract.add_run("Аннотация—")
    set_font(r, size=9, bold=True, italic=True)
    r = abstract.add_run(
        "Исследуется влияние температурного параметра β на геометрию пакетного градиента и низкоранговую адаптацию reward-модели. Эксперимент выполнен на 2000 реальных парах Anthropic HH-RLHF. В линейной постановке при изменении β от 0,1 до 10 угол градиента достиг 34,71°, а ESS/n уменьшился до 0,7141. В BERT-tiny геометрический эффект воспроизведён в пространстве LoRA-параметров. Внешний CPU-пилот на TinyLlama 1,1B обучал 0,0546% параметров и оценивался на стратифицированной выборке 92 примеров RewardBench. Micro accuracy изменилась с 42,39% до 45,65%, но 95%-й bootstrap-интервал изменения включал ноль (−4,35…10,87 п.п.; p=0,581); на 48 отложенных HH-парах точность снизилась с 47,92% до 45,83%. Результат подтверждает вычислительную применимость LoRA к модели миллиардного масштаба, но показывает необходимость категориальной внешней проверки."
    )
    set_font(r, size=9, bold=False, italic=False)

    keywords = doc.add_paragraph(style="Keywords")
    keywords.paragraph_format.first_line_indent = Inches(0)
    set_font(keywords.add_run("Ключевые слова—"), size=9, bold=True, italic=True)
    set_font(keywords.add_run("оптимизация предпочтений; геометрия градиента; LoRA; reward-модель; RewardBench; TinyLlama"), size=9, bold=False, italic=True)

    add_heading(doc, "ВВЕДЕНИЕ", 1)
    add_body(doc, "Обучение по человеческим предпочтениям стало основным инструментом согласования поведения языковых моделей с намерениями пользователя [1], [2]. Direct Preference Optimization (DPO) сводит этап оптимизации политики к логистической классификации пар предпочтительный–отклонённый ответ [3]. Более поздние алгоритмы, включая SimPO [4] и ORPO [5], меняют способ построения оценки, нормализацию длины или связь с этапом supervised fine-tuning. Несмотря на различия, во многих методах присутствует температурный параметр β, который масштабирует предпочтительный отступ.")
    add_body(doc, "Для одной пары множитель β не меняет луч градиента: он только умножает градиент отступа на положительный скаляр. Однако реальный шаг оптимизатора строится по мини-пакету непараллельных вкладов. Изменение β перераспределяет их веса и поэтому способно повернуть суммарный градиент. Одновременно парная потеря зависит только от разности оценок, тогда как поэлементная потеря отдельно закрепляет оценки предпочтительного и отклонённого ответов. Это различие связано с классической моделью Брэдли–Терри [6] и парным обучением ранжированию RankNet [7].")
    add_body(doc, "Цель работы — проверить оба эффекта на реальных данных человеческих предпочтений в полностью воспроизводимой постановке. В отличие от синтетических проверок исходных статей, здесь используется выборка HH-RLHF, публикуются все конфигурации, результаты отдельных запусков, контрольные суммы и тесты производных. Основные вопросы: насколько β поворачивает пакетный градиент; как меняется концентрация весов; различаются ли парная и поэлементная потери при одинаковом представлении и бюджете оптимизации.")
    add_body(doc, "Параметрически эффективное обучение дополняет эту постановку. LoRA ограничивает обновление весовой матрицы низкоранговым подпространством [13], QLoRA сочетает его с квантизацией [14], а LoRA+, DoRA и PiSSA уточняют скорости обучения, разложение веса и инициализацию адаптера [15]–[17]. Современная оценка reward-моделей показывает необходимость внешних и устойчивых проверок [18], а SePO переносит предпочтительный сигнал на информативные токены [19]. Поэтому второй эксперимент проверяет нашу геометрическую гипотезу непосредственно в пространстве LoRA-параметров.")
    add_body(doc, "Практическая актуальность связана с системами, где требуется выбрать лучший из нескольких сгенерированных ответов: контакт-центрами, финансовым комплаенсом, образовательной обратной связью, информационными материалами здравоохранения, промышленными инструкциями, государственными услугами и проверкой программного кода. В таких сценариях reward-модель должна быть вспомогательным фильтром, а не автономным источником юридически, клинически или технологически значимого решения. Межотраслевой профиль NIST требует контекстной оценки рисков генеративного ИИ [22], а OECD фиксирует одновременно расширение практических применений и риски непрозрачности, смещений и чрезмерного доверия [23], [24].")

    add_heading(doc, "МЕТОД", 1)
    add_heading(doc, "Функции потерь и геометрия", 2)
    add_body(doc, "Для пары ответов зададим xᵢ=φ(yᵢʷ)−φ(yᵢˡ) и линейный скорер sθ(y)=θᵀφ(y). Парная логистическая потеря имеет вид")
    add_native_equation(doc, "L_pair(θ;β)=n⁻¹∑ᵢ log(1+exp(−βθᵀxᵢ)).", 1)
    add_body(doc, "Её градиент является взвешенной суммой векторов xᵢ:")
    add_native_equation(doc, "∇L_pair=−βn⁻¹∑ᵢ σ(−βθᵀxᵢ)xᵢ.", 2)
    add_body(doc, "Веса wᵢ(β)=σ(−βθᵀxᵢ) зависят от величины и знака отступа. Поэтому при непараллельных xᵢ изменение β в общем случае меняет направление суммы. Концентрация весов оценивалась нормированным эффективным размером выборки")
    add_native_equation(doc, "ESS_norm=(∑ᵢwᵢ)²/(n∑ᵢwᵢ²).", 3)
    add_body(doc, "Поэлементная потеря отдельно классифицирует оба ответа:")
    add_native_equation(doc, "L_point=n⁻¹∑ᵢ[softplus(−βsθ(yᵢʷ))+softplus(βsθ(yᵢˡ))].", 4)
    add_body(doc, "Парная постановка инвариантна к общему сдвигу двух оценок, а поэлементная — нет. Следовательно, даже при совпадении порядка ответов функции потерь задают разную кривизну и разные ограничения на параметр θ.")

    add_heading(doc, "Данные и представление", 2)
    add_body(doc, "Использован открытый набор Anthropic HH-RLHF [2], [11]. Из обучающей части на 160800 строк детерминированно выбраны непересекающиеся страницы API; после исключения пар с полностью совпадающими chosen и rejected оставлены 2000 валидных пар. Выборка зафиксирована seed 20260918. SHA-256 канонического JSONL: e7a1ba050309838772fc8eca5506f601b91b556b7ab3fcddd3f77ee591e096bd.")
    add_body(doc, "Тексты преобразованы в signed feature hashing размерности 4096 по словным униграммам и биграммам с L2-нормировкой. Этот приём не требует словаря и воспроизводимо ограничивает память [8]. Представление намеренно фиксировано: измеряется геометрия функций предпочтения, а не качество конкретного трансформера.")
    add_body(doc, "Данные разделены на 1400 обучающих, 300 валидационных и 300 тестовых пар без пересечения. Для β∈{0,1; 0,3; 1; 3; 10} обе модели обучались Adam [9] в течение 800 шагов, batch size 128, learning rate 0,03, L2=10⁻³, seeds 11, 29, 47, 83 и 101. Для геометрической диагностики построена опорная парная модель при β=1 за 1000 шагов; градиенты вычислялись на неизменной валидационной части.")

    add_heading(doc, "Трансформер и LoRA", 2)
    add_body(doc, "Во втором эксперименте использован предобученный англоязычный BERT-tiny с двумя Transformer-слоями и скрытой размерностью 128. Тексты усекались слева до 192 токенов, чтобы сохранять конец ответа. Скалярная reward-голова применялась к представлению токена [CLS]. Контрольная модель обучала только голову; в LoRA-модели дополнительно обучались матрицы query и value всех attention-слоёв. Для исходной матрицы W низкоранговое обновление задавалось как")
    add_native_equation(doc, "W′=W+(α/r)BA,   rank(BA)≤r.", 5)
    add_body(doc, "Использованы r=8, α=16, dropout 0,05, β=0,3, AdamW, learning rate 5·10⁻⁴ для адаптеров и 10⁻³ для головы, weight decay 0,01, batch size 16 и seeds 11, 29, 47. Выполнялись четыре эпохи; checkpoint выбирался исключительно по минимальной валидационной парной потере. Тестовая часть применялась один раз после выбора. В LoRA обучались 8321 из 4394241 параметров. Для лучшего по validation loss запуска градиенты всех обучаемых параметров вычислялись на фиксированном пакете из 64 валидационных пар при пяти значениях β.")

    add_heading(doc, "Модель 1,1B и внешняя проверка", 2)
    add_body(doc, "Для проверки масштаба 1–3 млрд параметров использована TinyLlama-1.1B-Chat-v1.0 [21] как Llama-подобная reward-модель со скалярной классификационной головой. LoRA ранга 4 (α=8, dropout 0,05) применялась к q_proj и v_proj; обучались 565248 из 1035079680 параметров (0,0546%). Из-за CPU-ограничения выполнен заранее фиксированный пилот: 24 шага по 24 обучающим HH-парам, β=0,3, длина 64 токена. До и после адаптации модель проверена на 48 непересекающихся HH-парах и на 92 примерах RewardBench — по четыре примера из каждого из 23 поднаборов. Это диагностическая стратифицированная выборка, а не официальный агрегат leaderboard [18].")

    add_heading(doc, "Проверки и метрики", 2)
    add_body(doc, "Аналитические градиенты обеих потерь проверены центральными конечными разностями: абсолютная ошибка направленной производной составила 6,65·10⁻¹³ для парной и 8,11·10⁻¹¹ для поэлементной потери. Проверены границы ESS, углы опорных векторов, отсутствие пересечения выборок, повторяемость хеширования и контрольные суммы данных. Оценивались угол, норма, ESS, насыщение, ranking accuracy, log loss, Brier и ECE [10]. Для точности рассчитаны bootstrap-интервалы [12]; изменение на одинаковых RewardBench-примерах дополнительно проверено двусторонним точным sign test.")

    add_heading(doc, "РЕЗУЛЬТАТЫ", 1)
    geometry = data["geometry"]
    add_table(
        doc,
        "ГЕОМЕТРИЯ ГРАДИЕНТА НА ВАЛИДАЦИОННОЙ ВЫБОРКЕ",
        ["β", "Угол, °", "ESS/n", "Насыщение"],
        [[f"{r['beta']:g}", f"{r['angle_to_beta_1_deg']:.2f}", f"{r['normalized_ess']:.4f}", f"{r['saturated_weight_share']:.4f}"] for r in geometry],
        [650, 1300, 1300, 1550],
    )
    add_body(doc, "Таблица I и рис. 1 показывают немонотонный поворот относительно опорного β=1. При β=0,1 угол равен 17,22°, при β=3 — 20,89°, а при β=10 — 34,71°. Одновременно ESS/n снижается с 0,9998 до 0,7141. При β=10 уже 21,67% пар имеют насыщенные веса. Таким образом, заменить β изменением learning rate нельзя: направление и фактический состав мини-пакета меняются вместе с масштабом.")
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = p.add_run().add_picture(str(FIGURE), width=Inches(3.25))
    picture._inline.docPr.set("descr", "Два графика зависимости угла градиента и нормированного эффективного размера выборки от beta")
    picture._inline.docPr.set("title", "Геометрия пакетного градиента на HH-RLHF")
    cap = doc.add_paragraph(style="figure caption")
    cap.paragraph_format.keep_together = True
    set_font(cap.add_run("Поворот пакетного градиента и уменьшение нормированного ESS при изменении β."), size=8)

    summary = data["summary"]
    add_table(
        doc,
        "ТЕСТОВЫЕ РЕЗУЛЬТАТЫ ПЯТИ ЗАПУСКОВ",
        ["Потеря", "β", "Точность", "Log loss", "Brier"],
        [["парная" if r["method"] == "pairwise" else "поэлем.", f"{r['beta']:g}", f"{r['accuracy_mean']:.3f}±{r['accuracy_sd']:.3f}", fmt3(r["pair_loss_mean"]), fmt3(r["brier_mean"])] for r in summary],
        [1100, 500, 1350, 950, 900],
    )
    add_body(doc, "Лучший результат обеих моделей получен при β=0,3. Парная потеря достигла точности 0,5627±0,0072, поэлементная — 0,5433±0,0129. Среднее преимущество парной модели по совпадающим seed составило 1,93 процентного пункта, но вывод ограничен одной выборкой и одним классом признаков. При β=10 точность снизилась до 0,4800 и 0,4833; логистическая потеря выросла до 1,069 и 1,218 соответственно. Следовательно, рост нормы градиента при большом β не преобразуется в улучшение ранжирования.")
    add_body(doc, "Вероятностная калибровка во всех режимах остаётся слабой: минимальный Brier score равен 0,2484 для парной модели и 0,2492 для поэлементной. Это ожидаемо для компактного линейного скорера и подчёркивает, что ranking accuracy и уверенность следует анализировать раздельно.")

    add_heading(doc, "РЕЗУЛЬТАТЫ LORA-АДАПТАЦИИ", 1)
    lora_summary = lora_data["summary"]
    add_table(
        doc,
        "СРАВНЕНИЕ ТРАНСФОРМЕРНЫХ REWARD-МОДЕЛЕЙ",
        ["Метод", "Обуч. парам.", "Доля, %", "Точность", "Log loss"],
        [[
            "только head" if r["method"] == "head_only" else "LoRA r=8",
            str(r["trainable_parameters"]),
            f"{100*r['trainable_share']:.3f}",
            f"{r['accuracy_mean']:.3f}±{r['accuracy_sd']:.3f}",
            f"{r['pair_loss_mean']:.3f}",
        ] for r in lora_summary],
        [1150, 1000, 800, 1450, 900],
    )
    add_body(doc, "LoRA не повысила среднюю ranking accuracy: 0,5333±0,0200 против 0,5467±0,0058 у контроля. Средняя парная потеря также практически совпала: 0,6872 и 0,6865. Все три LoRA-запуска выбрали первую эпоху по validation loss, тогда как контроль в среднем выбрал эпоху 1,67. Это свидетельствует о раннем переобучении адаптера и не позволяет заявлять преимущество LoRA на 1400 обучающих парах. Отдельный лучший LoRA-seed достиг 0,5533, но единичный результат не использовался как основной вывод.")
    add_table(
        doc,
        "ГЕОМЕТРИЯ ГРАДИЕНТА В ПРОСТРАНСТВЕ LORA",
        ["β", "Угол, °", "ESS/n", "Насыщение"],
        [[f"{r['beta']:g}", f"{r['angle_to_beta_1_deg']:.2f}", f"{r['normalized_ess']:.4f}", f"{r['saturated_weight_share']:.4f}"] for r in lora_data["geometry"]],
        [650, 1300, 1300, 1550],
    )
    add_body(doc, "Геометрический эффект воспроизвёлся в обучаемом нейросетевом подпространстве. При β=0,1 угол к градиенту β=1 равен 22,99°, при β=3 — 19,44°, при β=10 — 30,39°. Одновременно ESS/n уменьшился до 0,6199, а доля насыщенных весов выросла до 0,3125. Следовательно, температурное масштабирование меняет направление обновления не только у линейного суррогата, но и у LoRA-адаптера.")

    add_heading(doc, "ВНЕШНИЙ ПИЛОТ TINYLLAMA 1,1B", 1)
    bert_rb = {r["category"]: r["accuracy"] for r in external["bert_rewardbench"]}
    tiny_before = {r["category"]: r["accuracy"] for r in external["tinyllama_rewardbench_before"]}
    tiny_after = {r["category"]: r["accuracy"] for r in external["tinyllama_rewardbench_after"]}
    rb_categories = ["Chat", "Chat Hard", "Safety", "Reasoning", "Macro categories", "Micro examples"]
    rb_labels = {
        "Chat": "Chat",
        "Chat Hard": "Chat Hard",
        "Safety": "Safety",
        "Reasoning": "Reasoning",
        "Macro categories": "Macro",
        "Micro examples": "Micro",
    }
    add_table(
        doc,
        "ТОЧНОСТЬ НА СТРАТИФИЦИРОВАННОЙ ВЫБОРКЕ REWARDBENCH, %",
        ["Категория", "BERT LoRA", "1,1B до", "1,1B после"],
        [[rb_labels[c], f"{100*bert_rb[c]:.1f}", f"{100*tiny_before[c]:.1f}", f"{100*tiny_after[c]:.1f}"] for c in rb_categories],
        [1500, 1100, 1100, 1100],
    )
    change = external["tinyllama_rewardbench_paired_change"]
    add_body(doc, "После 24 LoRA-шагов micro accuracy TinyLlama на 92 внешних примерах изменилась с 42,39% до 45,65% (+3,26 п.п.): улучшились 8 примеров, ухудшились 5, без изменения остались 79. Однако 95%-й bootstrap-интервал парного изменения равен −4,35…10,87 п.п., а двусторонний exact sign test даёт p=0,581. На 48 отложенных HH-парах точность снизилась с 47,92% до 45,83%. По категориям выросли Chat, Chat Hard и Safety, но Reasoning снизилась с 50,0% до 46,4%. Следовательно, пилот демонстрирует воспроизводимый протокол переноса, но не статистически подтверждённое преимущество.")

    add_heading(doc, "ОБСУЖДЕНИЕ", 1)
    add_body(doc, "Наблюдаемая связь между углом и ESS согласуется с теоретическим механизмом: увеличение β подавляет вклад пар с положительным отступом и концентрирует обновление на трудных или ошибочных примерах. При малом β веса почти равны, но направление также отличается от β=1, поскольку опорная модель уже создаёт неоднородные отступы. Симметрия углов относительно β=1 не ожидается: логистическое перевзвешивание нелинейно.")
    add_body(doc, "Сравнение потерь показывает, что различие между разностной и абсолютной идентификацией сохраняется на реальных парах. LoRA-эксперимент переносит анализ с фиксированных признаков на внутренние представления Transformer, но отсутствие среднего прироста accuracy ограничивает практический вывод. Результат согласуется с идеей KTO о зависимости подходящей функции потерь от предположений о человеческой полезности [20]: параметризация адаптера сама по себе не устраняет несоответствие малой выборки и целевой метрики.")
    add_body(doc, "Секторное применение требует оценивать именно критический тип ошибки. Для поддержки и электронной торговли это фактическая полезность и ошибочная эскалация; для финансов — комплаенс и конфиденциальность; для образования — правильность рассуждения; для здравоохранения и промышленности — безопасный отказ и обязательный контроль специалиста; для государства — трассируемость, равное обращение и защита данных. Разнонаправленный результат RewardBench показывает, почему средний балл нельзя использовать как единственный допуск к эксплуатации.")
    add_body(doc, "Ограничения: один англоязычный обучающий набор; только три BERT-seed; один seed, 24 пары и короткие последовательности в CPU-пилоте 1,1B; 92 внешних примера вместо полного RewardBench; отсутствие отраслевой разметки. Дальнейшая работа должна выполнить полный benchmark, увеличить число seed и пар, сравнить LoRA/LoRA+/PiSSA на 1–3B моделях, добавить M-RewardBench и русскоязычный корпус, а также предварительно зарегистрировать критерии успеха по каждой категории.")

    add_heading(doc, "ЗАКЛЮЧЕНИЕ", 1)
    add_body(doc, "Эксперимент подтвердил, что β является геометрическим гиперпараметром: угол достиг 34,71° в линейной модели и 30,39° в пространстве LoRA. Адаптация BERT-tiny не дала среднего преимущества, а CPU-пилот TinyLlama 1,1B изменил внешнюю micro accuracy на +3,26 п.п. без статистической значимости и ухудшил отложенную HH-точность. Поэтому масштаб модели и малая доля обучаемых параметров не заменяют внешнюю проверку. Для практики β, функцию потерь, LoRA-конфигурацию и отраслевой критерий допуска следует выбирать совместно, сохраняя категориальные метрики и human-in-the-loop контроль.")

    add_heading(doc, "ДОСТУПНОСТЬ ДАННЫХ И КОДА", 5)
    add_body(doc, "Код, конфигурации, результаты каждого запуска, контрольные суммы и русская версия статьи опубликованы по адресу https://github.com/PetrNikitin20/optimML. Исходные пары загружаются из Anthropic HH-RLHF; сырые тексты не включены в репозиторий.")

    add_heading(doc, "СПИСОК ЛИТЕРАТУРЫ", 5)
    references = [
        "L. Ouyang et al., “Training language models to follow instructions with human feedback,” in Advances in Neural Information Processing Systems, vol. 35, pp. 27730–27744, 2022.",
        "Y. Bai et al., “Training a helpful and harmless assistant with reinforcement learning from human feedback,” arXiv:2204.05862, 2022.",
        "R. Rafailov, A. Sharma, E. Mitchell, C. D. Manning, S. Ermon, and C. Finn, “Direct preference optimization: Your language model is secretly a reward model,” in Advances in Neural Information Processing Systems, vol. 36, 2023, pp. 53728–53741.",
        "Y. Meng, M. Xia, and D. Chen, “SimPO: Simple preference optimization with a reference-free reward,” in Advances in Neural Information Processing Systems, vol. 37, 2024.",
        "J. Hong, N. Lee, and J. Thorne, “ORPO: Monolithic preference optimization without reference model,” in Proc. EMNLP, 2024, pp. 11170–11189, doi: 10.18653/v1/2024.emnlp-main.626.",
        "R. A. Bradley and M. E. Terry, “Rank analysis of incomplete block designs: I. The method of paired comparisons,” Biometrika, vol. 39, no. 3/4, pp. 324–345, 1952, doi: 10.1093/biomet/39.3-4.324.",
        "C. J. C. Burges et al., “Learning to rank using gradient descent,” in Proc. 22nd Int. Conf. Machine Learning, 2005, pp. 89–96, doi: 10.1145/1102351.1102363.",
        "K. Weinberger, A. Dasgupta, J. Langford, A. Smola, and J. Attenberg, “Feature hashing for large scale multitask learning,” in Proc. 26th Int. Conf. Machine Learning, 2009, pp. 1113–1120, doi: 10.1145/1553374.1553516.",
        "D. P. Kingma and J. Ba, “Adam: A method for stochastic optimization,” in Proc. 3rd Int. Conf. Learning Representations, 2015.",
        "C. Guo, G. Pleiss, Y. Sun, and K. Q. Weinberger, “On calibration of modern neural networks,” in Proc. 34th Int. Conf. Machine Learning, PMLR 70, 2017, pp. 1321–1330.",
        "Anthropic, “HH-RLHF dataset card,” Hugging Face, 2022. [Online]. Available: https://huggingface.co/datasets/Anthropic/hh-rlhf. Accessed: Sep. 18, 2026.",
        "B. Efron and R. J. Tibshirani, An Introduction to the Bootstrap. New York, NY, USA: Chapman & Hall, 1993.",
        "E. J. Hu et al., “LoRA: Low-rank adaptation of large language models,” in Proc. Int. Conf. Learning Representations, 2022.",
        "T. Dettmers, A. Pagnoni, A. Holtzman, and L. Zettlemoyer, “QLoRA: Efficient finetuning of quantized LLMs,” in Advances in Neural Information Processing Systems, vol. 36, 2023, pp. 10088–10115.",
        "S. Hayou, N. Ghosh, and B. Yu, “LoRA+: Efficient low rank adaptation of large models,” in Proc. 41st Int. Conf. Machine Learning, PMLR 235, 2024, pp. 17783–17806.",
        "S.-Y. Liu et al., “DoRA: Weight-decomposed low-rank adaptation,” in Proc. 41st Int. Conf. Machine Learning, PMLR 235, 2024, pp. 32100–32121.",
        "F. Meng, Z. Wang, and M. Zhang, “PiSSA: Principal singular values and singular vectors adaptation of large language models,” in Advances in Neural Information Processing Systems, vol. 37, 2024, doi: 10.52202/079017-3846.",
        "N. Lambert et al., “RewardBench: Evaluating reward models for language modeling,” in Findings of NAACL, 2025, pp. 1755–1797, doi: 10.18653/v1/2025.findings-naacl.96.",
        "K. Yang, Z. Liu, Q. Xie, J. Huang, E. Min, and S. Ananiadou, “Selective preference optimization via token-level reward function estimation,” in Proc. EMNLP, 2025, pp. 7032–7056, doi: 10.18653/v1/2025.emnlp-main.359.",
        "K. Ethayarajh, W. Xu, N. Muennighoff, D. Jurafsky, and D. Kiela, “Model alignment as prospect theoretic optimization,” in Proc. 41st Int. Conf. Machine Learning, PMLR 235, 2024, pp. 12634–12651.",
        "P. Zhang, G. Zeng, T. Wang, and W. Lu, “TinyLlama: An open-source small language model,” arXiv:2401.02385, 2024.",
        "C. Autio et al., Artificial Intelligence Risk Management Framework: Generative Artificial Intelligence Profile, NIST AI 600-1, 2024, doi: 10.6028/NIST.AI.600-1.",
        "OECD, Governing with Artificial Intelligence: The State of Play and Way Forward in Core Government Functions. Paris, France: OECD Publishing, 2025, doi: 10.1787/795de142-en.",
        "OECD, Generative AI and the SME Workforce. Paris, France: OECD Publishing, 2025.",
    ]
    for i, reference in enumerate(references, 1):
        add_reference(doc, i, reference)

    balancing_section = doc.add_section(WD_SECTION.CONTINUOUS)
    balancing_section.left_margin = Inches(0.62)
    balancing_section.right_margin = Inches(0.62)
    balancing_section.top_margin = Inches(0.75)
    balancing_section.bottom_margin = Inches(1.0)
    set_columns(balancing_section, 1, 360)

    doc.core_properties.title = "Геометрия градиента и LoRA-адаптация при обучении по предпочтениям"
    doc.core_properties.author = "Алексей П. Малахов"
    doc.core_properties.subject = "Русский проект статьи для SUMMA 2026"
    doc.core_properties.keywords = "preference optimization, gradient geometry, LoRA, reward model, HH-RLHF"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
